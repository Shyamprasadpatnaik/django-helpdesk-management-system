import csv
import io
from datetime import datetime, timedelta

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Count, Avg, F, Q, DurationField
from django.utils import timezone

from tickets.models import Ticket, TicketPriority, TicketStatus
from departments.models import Department
from accounts.models import User, Role
from accounts.decorators import hr_required, super_admin_required

def _filter_tickets(request):
    """Shared filter helper for report dashboard and export actions."""
    qs = Ticket.objects.select_related('department', 'priority', 'created_by', 'assigned_to').all()

    dept_id = request.GET.get('department')
    priority_id = request.GET.get('priority')
    status = request.GET.get('status')
    tech_id = request.GET.get('technician')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    if dept_id:
        qs = qs.filter(department_id=dept_id)
    if priority_id:
        qs = qs.filter(priority_id=priority_id)
    if status:
        qs = qs.filter(status=status)
    if tech_id:
        qs = qs.filter(assigned_to_id=tech_id)
    if date_from:
        try:
            df = datetime.strptime(date_from, '%Y-%m-%d')
            qs = qs.filter(created_at__gte=df)
        except ValueError:
            pass
    if date_to:
        try:
            dt = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1)
            qs = qs.filter(created_at__lt=dt)
        except ValueError:
            pass

    return qs


@login_required
@hr_required
def report_dashboard(request):
    """
    Analytics & Reporting dashboard with interactive filters and charts.
    """
    tickets = _filter_tickets(request)

    total_count = tickets.count()
    resolved_count = tickets.filter(status__in=[TicketStatus.RESOLVED, TicketStatus.CLOSED]).count()
    breached_count = tickets.filter(sla_breached=True).count()

    sla_compliance_rate = 100
    if total_count > 0:
        sla_compliance_rate = round(((total_count - breached_count) / total_count) * 100, 1)

    # Department performance
    dept_stats = Department.objects.filter(is_active=True).annotate(
        total=Count('tickets'),
        open=Count('tickets', filter=~Q(tickets__status__in=[TicketStatus.RESOLVED, TicketStatus.CLOSED])),
        breached=Count('tickets', filter=Q(tickets__sla_breached=True))
    )

    # Common issue categories
    common_categories = tickets.values('category__name').annotate(
        count=Count('id')
    ).filter(category__name__isnull=False).order_by('-count')[:8]

    # Technician workload summary
    tech_stats = User.objects.filter(role=Role.TECHNICIAN, is_active=True).annotate(
        assigned_total=Count('assigned_tickets'),
        resolved_total=Count('assigned_tickets', filter=Q(assigned_tickets__status__in=[TicketStatus.RESOLVED, TicketStatus.CLOSED]))
    )

    departments = Department.objects.filter(is_active=True)
    priorities = TicketPriority.objects.all().order_by('order')
    technicians = User.objects.filter(role=Role.TECHNICIAN, is_active=True)

    context = {
        'total_count': total_count,
        'resolved_count': resolved_count,
        'breached_count': breached_count,
        'sla_compliance_rate': sla_compliance_rate,
        'dept_stats': dept_stats,
        'common_categories': common_categories,
        'tech_stats': tech_stats,
        'departments': departments,
        'priorities': priorities,
        'technicians': technicians,
        'statuses': TicketStatus.choices,
        'selected_dept': request.GET.get('department', ''),
        'selected_priority': request.GET.get('priority', ''),
        'selected_status': request.GET.get('status', ''),
        'selected_tech': request.GET.get('technician', ''),
        'date_from': request.GET.get('date_from', ''),
        'date_to': request.GET.get('date_to', ''),
    }
    return render(request, 'reports/report_dashboard.html', context)


@login_required
@hr_required
def export_csv(request):
    """Exports filtered tickets to CSV format."""
    tickets = _filter_tickets(request)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="tickets_export_{timezone.now().strftime("%Y%m%d_%H%M")}.csv"'

    writer = csv.writer(response)
    writer.writerow([
        'Ticket Number', 'Title', 'Department', 'Category', 'Priority',
        'Status', 'Created By', 'Assigned Technician', 'Equipment/Machine',
        'Location', 'Created At', 'Due SLA At', 'SLA Breached', 'Reopened Count'
    ])

    for t in tickets:
        writer.writerow([
            t.ticket_number,
            t.title,
            t.department.name if t.department else '',
            t.category.name if t.category else 'General',
            t.priority.name if t.priority else '',
            t.get_status_display(),
            t.created_by.display_name if t.created_by else '',
            t.assigned_to.display_name if t.assigned_to else 'Unassigned',
            t.equipment_number or '',
            f"{t.building or ''} {t.room_area or ''}".strip(),
            t.created_at.strftime('%Y-%m-%d %H:%M'),
            t.due_at.strftime('%Y-%m-%d %H:%M') if t.due_at else '',
            'Yes' if t.sla_breached else 'No',
            t.reopened_count
        ])

    return response


@login_required
@hr_required
def export_excel(request):
    """Exports filtered tickets to formatted Microsoft Excel (.xlsx) using openpyxl."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    tickets = _filter_tickets(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "JK Paper Tickets Report"

    # Title Banner
    ws.merge_cells('A1:L1')
    title_cell = ws['A1']
    title_cell.value = "JK Paper Mill - Ticket Management System Report"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Timestamp row
    ws['A2'] = f"Generated on: {timezone.now().strftime('%Y-%m-%d %H:%M')} | Total Records: {tickets.count()}"
    ws['A2'].font = Font(name="Arial", size=10, italic=True)

    headers = [
        'Ticket ID', 'Title', 'Department', 'Category', 'Priority',
        'Status', 'Reporter', 'Technician', 'Equipment Tag',
        'Created Date', 'SLA Due', 'SLA Breached'
    ]

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header_title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 24

    for row_idx, t in enumerate(tickets, 5):
        row_data = [
            t.ticket_number,
            t.title,
            t.department.name if t.department else '',
            t.category.name if t.category else 'General',
            t.priority.name if t.priority else '',
            t.get_status_display(),
            t.created_by.display_name if t.created_by else '',
            t.assigned_to.display_name if t.assigned_to else 'Unassigned',
            t.equipment_number or '',
            t.created_at.strftime('%Y-%m-%d %H:%M'),
            t.due_at.strftime('%Y-%m-%d %H:%M') if t.due_at else '',
            'Yes' if t.sla_breached else 'No'
        ]
        for col_idx, value in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.font = Font(name="Arial", size=10)
            c.border = thin_border
            # Zebra striping
            if row_idx % 2 == 0:
                c.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="jkpaper_tickets_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    return response


@login_required
@hr_required
def export_pdf(request):
    """Exports summary report to PDF using reportlab."""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    tickets = _filter_tickets(request)[:100]  # Limit to 100 for clean pagination

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20
    )
    elements = []
    styles = getSampleStyleSheet()

    # Custom Header Style
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=4
    )
    sub_style = ParagraphStyle(
        'SubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=15
    )

    elements.append(Paragraph("JK Paper Mill - Enterprise Ticket Management Report", title_style))
    elements.append(Paragraph(f"Generated on {timezone.now().strftime('%d %B %Y, %H:%M')} | Scope: {tickets.count()} Tickets", sub_style))

    # Table Header & Data
    table_data = [
        ['Ticket ID', 'Title', 'Department', 'Priority', 'Status', 'Technician', 'Created At', 'SLA Breached']
    ]

    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=10)

    for t in tickets:
        table_data.append([
            t.ticket_number,
            Paragraph(t.title[:35] + ('...' if len(t.title) > 35 else ''), cell_style),
            t.department.code if t.department else '',
            t.priority.name if t.priority else '',
            t.get_status_display(),
            t.assigned_to.display_name[:20] if t.assigned_to else 'Unassigned',
            t.created_at.strftime('%d/%m/%y %H:%M'),
            'YES' if t.sla_breached else 'NO'
        ])

    pdf_table = Table(table_data, colWidths=[80, 190, 75, 60, 95, 110, 85, 65])
    pdf_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
    ]))

    elements.append(pdf_table)
    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="jkpaper_tickets_{timezone.now().strftime("%Y%m%d_%H%M")}.pdf"'
    return response
